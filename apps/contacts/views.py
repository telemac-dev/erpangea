import csv
import io
import openpyxl
from django.shortcuts import render, get_object_or_404, redirect
from django.views import View
from django.views.generic import ListView, DetailView, CreateView, UpdateView
from django.contrib.auth.mixins import LoginRequiredMixin
from django.contrib import messages
from django.http import HttpResponse, JsonResponse
from django.urls import reverse_lazy, reverse
from django.db.models import Q, Count
from django.utils.translation import gettext_lazy as _
from django.core.exceptions import ValidationError

from .models import Contact, ContactTag, ContactTypeChoices, AddressTypeChoices, DocTypeChoices
from .forms import (
    ContactForm,
    SubordinateContactForm,
    ContactTagForm,
    ContactMergeForm,
    ContactUploadImportForm
)
from .validators import clean_doc_digits, validate_cpf, validate_cnpj, format_document
from apps.accounts.permissions import RoleRequiredMixin

class ContactListView(LoginRequiredMixin, ListView):
    model = Contact
    paginate_by = 24
    context_object_name = 'contacts'

    def get_template_names(self):
        if self.request.htmx and self.request.GET.get('page_only'):
            view_mode = self.request.GET.get('view', 'kanban')
            return [f'contacts/partials/{view_mode}_content.html']
        return ['contacts/list.html']

    def get_queryset(self):
        qs = Contact.objects.select_related('parent', 'salesperson').prefetch_related('tags')
        
        # 1. Filtro de Arquivados vs Ativos
        filter_type = self.request.GET.get('filter', '')
        if filter_type == 'archived':
            qs = qs.filter(is_active=False)
        else:
            qs = qs.filter(is_active=True)

        # 2. Filtros pré-definidos
        if filter_type == 'company':
            qs = qs.filter(contact_type=ContactTypeChoices.COMPANY)
        elif filter_type == 'individual':
            qs = qs.filter(contact_type=ContactTypeChoices.INDIVIDUAL)

        # 3. Busca textual
        q = self.request.GET.get('q', '').strip()
        if q:
            clean_q_digits = clean_doc_digits(q)
            query = Q(name__icontains=q) | Q(trade_name__icontains=q) | Q(email__icontains=q) | Q(city__icontains=q) | Q(tags__name__icontains=q)
            if clean_q_digits:
                query |= Q(doc_number__icontains=clean_q_digits)
            qs = qs.filter(query).distinct()

        # 4. Agrupamento (Group By)
        group_by = self.request.GET.get('group_by', '')
        if group_by:
            if group_by == 'city':
                qs = qs.order_by('city', 'name')
            elif group_by == 'state':
                qs = qs.order_by('state', 'name')
            elif group_by == 'salesperson':
                qs = qs.order_by('salesperson__email', 'name')
            elif group_by == 'parent':
                qs = qs.order_by('parent__name', 'name')
        else:
            ordering = self.request.GET.get('ordering', 'name')
            qs = qs.order_by(ordering)

        return qs

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        view_mode = self.request.GET.get('view', 'kanban')
        context['view_mode'] = view_mode
        context['current_filter'] = self.request.GET.get('filter', '')
        context['current_group'] = self.request.GET.get('group_by', '')
        context['query'] = self.request.GET.get('q', '')
        context['total_contacts'] = self.get_queryset().count()
        context['tags'] = ContactTag.objects.all()
        return context

class ContactDetailView(LoginRequiredMixin, DetailView):
    model = Contact
    template_name = 'contacts/detail.html'
    context_object_name = 'contact'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['subordinates'] = self.object.subordinates.filter(is_active=True).select_related('parent')
        return context

class ContactCreateView(LoginRequiredMixin, CreateView):
    model = Contact
    form_class = ContactForm
    template_name = 'contacts/form.html'
    success_url = reverse_lazy('contacts:list')

    def get_initial(self):
        initial = super().get_initial()
        c_type = self.request.GET.get('type')
        if c_type == 'individual':
            initial['contact_type'] = ContactTypeChoices.INDIVIDUAL
            initial['doc_type'] = DocTypeChoices.CPF
        else:
            initial['contact_type'] = ContactTypeChoices.COMPANY
            initial['doc_type'] = DocTypeChoices.CNPJ
        return initial

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['is_edit'] = False
        return context

    def form_valid(self, form):
        messages.success(self.request, _("Contato cadastrado com sucesso!"))
        return super().form_valid(form)

class ContactUpdateView(LoginRequiredMixin, UpdateView):
    model = Contact
    form_class = ContactForm
    template_name = 'contacts/form.html'

    def get_success_url(self):
        return reverse('contacts:detail', kwargs={'pk': self.object.pk})

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['is_edit'] = True
        return context

    def form_valid(self, form):
        messages.success(self.request, _("Contato atualizado com sucesso!"))
        return super().form_valid(form)

class ContactArchiveToggleView(LoginRequiredMixin, View):
    def post(self, request, pk):
        contact = get_object_or_404(Contact, pk=pk)
        if contact.is_active:
            contact.archive()
            messages.warning(request, _(f"Contato '{contact.name}' foi arquivado com sucesso."))
        else:
            contact.unarchive()
            messages.success(request, _(f"Contato '{contact.name}' foi desarquivado com sucesso."))
        return redirect('contacts:list')

class ContactValidateDocumentView(LoginRequiredMixin, View):
    """
    Endpoint assincrono para validacao em tempo real de CPF/CNPJ via HTMX.
    """
    def post(self, request):
        doc_number = request.POST.get('doc_number', '').strip()
        doc_type = request.POST.get('doc_type', 'CNPJ')
        contact_id = request.POST.get('contact_id', '').strip()

        if not doc_number:
            return HttpResponse("")

        digits = clean_doc_digits(doc_number)
        
        # Validacao matematica
        error_msg = None
        try:
            if doc_type == 'CPF':
                validate_cpf(digits)
            elif doc_type == 'CNPJ':
                validate_cnpj(digits)
        except ValidationError as e:
            error_msg = e.messages[0]

        if error_msg:
            return HttpResponse(
                f'<div class="text-danger small mt-1 d-flex align-items-center gap-1">'
                f'<i class="bi bi-x-circle-fill"></i> {error_msg}'
                f'</div>'
            )

        # Checagem de duplicidade entre contatos ativos
        qs = Contact.objects.filter(doc_number=digits, is_active=True)
        if contact_id:
            qs = qs.exclude(pk=contact_id)

        if qs.exists():
            existing = qs.first()
            return HttpResponse(
                f'<div class="text-warning small mt-1 d-flex align-items-center gap-1">'
                f'<i class="bi bi-exclamation-triangle-fill"></i> Atenção: Documento já cadastrado para "{existing.name}".'
                f'</div>'
            )

        formatted = format_document(digits, doc_type)
        return HttpResponse(
            f'<div class="text-success small mt-1 d-flex align-items-center gap-1">'
            f'<i class="bi bi-check-circle-fill"></i> {doc_type} Válido: <strong>{formatted}</strong>'
            f'</div>'
        )

class SubordinateContactCreateView(LoginRequiredMixin, View):
    """
    Modal dinâmico para cadastrar pessoas de contato subordinadas ou endereços de entrega/cobrança.
    """
    def get(self, request, parent_id):
        parent = get_object_or_404(Contact, pk=parent_id)
        form = SubordinateContactForm(initial={
            'city': parent.city,
            'state': parent.state,
            'country': parent.country,
        })
        return render(request, 'contacts/partials/subordinate_modal.html', {
            'form': form,
            'parent': parent
        })

    def post(self, request, parent_id):
        parent = get_object_or_404(Contact, pk=parent_id)
        form = SubordinateContactForm(request.POST)
        if form.is_valid():
            sub = form.save(commit=False)
            sub.parent = parent
            sub.contact_type = ContactTypeChoices.INDIVIDUAL
            sub.save()
            messages.success(request, _(f"Vínculo '{sub.name}' adicionado com sucesso!"))
            response = HttpResponse("")
            response['HX-Refresh'] = 'true'
            return response
        return render(request, 'contacts/partials/subordinate_modal.html', {
            'form': form,
            'parent': parent
        })

class ContactMergeView(LoginRequiredMixin, View):
    """
    Assistente de deduplicação e mesclagem de contatos.
    """
    def get(self, request):
        ids = request.GET.get('ids', '').split(',')
        contacts = Contact.objects.filter(id__in=[i for i in ids if i], is_active=True)
        if contacts.count() < 2:
            return HttpResponse("<div class='p-3 text-danger'>Selecione pelo menos 2 contatos para mesclar.</div>")
        form = ContactMergeForm(contacts=contacts)
        return render(request, 'contacts/partials/merge_modal.html', {
            'contacts': contacts,
            'form': form,
            'contact_ids': ','.join(str(c.pk) for c in contacts)
        })

    def post(self, request):
        ids = request.POST.get('contact_ids', '').split(',')
        contacts = Contact.objects.filter(id__in=[i for i in ids if i], is_active=True)
        form = ContactMergeForm(request.POST, contacts=contacts)
        
        if form.is_valid():
            dest = form.cleaned_data['destination_contact']
            duplicates = contacts.exclude(pk=dest.pk)

            # Redireciona contatos subordinados
            Contact.objects.filter(parent__in=duplicates).update(parent=dest)

            # Arquiva contatos duplicados consolidados
            for dup in duplicates:
                dup.archive()

            messages.success(request, _(f"Contatos mesclados com sucesso no cadastro de '{dest.name}'."))
            response = HttpResponse("")
            response['HX-Refresh'] = 'true'
            return response

        return render(request, 'contacts/partials/merge_modal.html', {
            'contacts': contacts,
            'form': form,
            'contact_ids': ','.join(str(c.pk) for c in contacts)
        })

class ContactExportView(LoginRequiredMixin, View):
    """
    Exportação de contatos para CSV ou XLSX.
    """
    def get(self, request):
        fmt = request.GET.get('format', 'csv')
        contacts = Contact.objects.filter(is_active=True).select_related('parent')

        if fmt == 'xlsx':
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Contatos ERPangea"
            ws.append([
                'ID', 'Tipo', 'Nome / Razão Social', 'Nome Fantasia', 'Empresa Vinculada',
                'Tipo Doc', 'Documento', 'Inscrição Estadual', 'E-mail', 'Telefone', 'Celular',
                'Logradouro', 'Número', 'Bairro', 'Cidade', 'Estado', 'CEP'
            ])
            for c in contacts:
                ws.append([
                    str(c.id), c.get_contact_type_display(), c.name, c.trade_name,
                    c.parent.name if c.parent else '', c.doc_type, c.formatted_doc_number,
                    c.state_registration, c.email, c.phone, c.mobile,
                    c.street, c.number, c.neighborhood, c.city, c.state, c.postal_code
                ])
            response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
            response['Content-Disposition'] = 'attachment; filename="contatos_erpangea.xlsx"'
            wb.save(response)
            return response
        else:
            response = HttpResponse(content_type='text/csv; charset=utf-8')
            response['Content-Disposition'] = 'attachment; filename="contatos_erpangea.csv"'
            writer = csv.writer(response)
            writer.writerow([
                'ID', 'Tipo', 'Nome / Razão Social', 'Nome Fantasia', 'Empresa Vinculada',
                'Tipo Doc', 'Documento', 'Inscrição Estadual', 'E-mail', 'Telefone', 'Celular',
                'Logradouro', 'Número', 'Bairro', 'Cidade', 'Estado', 'CEP'
            ])
            for c in contacts:
                writer.writerow([
                    str(c.id), c.get_contact_type_display(), c.name, c.trade_name,
                    c.parent.name if c.parent else '', c.doc_type, c.formatted_doc_number,
                    c.state_registration, c.email, c.phone, c.mobile,
                    c.street, c.number, c.neighborhood, c.city, c.state, c.postal_code
                ])
            return response

class ContactTemplateDownloadView(LoginRequiredMixin, View):
    """
    Download do modelo padrão de planilha para importação.
    """
    def get(self, request):
        response = HttpResponse(content_type='text/csv; charset=utf-8')
        response['Content-Disposition'] = 'attachment; filename="modelo_importacao_contatos.csv"'
        writer = csv.writer(response)
        writer.writerow([
            'tipo_contato', 'nome_razao_social', 'nome_fantasia', 'tipo_documento',
            'numero_documento', 'email', 'telefone', 'celular', 'cidade', 'estado', 'cep'
        ])
        writer.writerow([
            'COMPANY', 'Exemplo Engenharia e Fundações Ltda', 'Exemplo Geotecnia', 'CNPJ',
            '11.222.333/0001-81', 'contato@exemplogeo.com.br', '(11) 3333-5555', '(11) 9 9999-8888', 'São Paulo', 'SP', '01310-100'
        ])
        writer.writerow([
            'INDIVIDUAL', 'Carlos Eduardo da Silva', '', 'CPF',
            '123.456.789-00', 'carlos.silva@exemplo.com.br', '', '(11) 9 8888-7777', 'São Paulo', 'SP', '01310-100'
        ])
        return response

class ContactImportView(LoginRequiredMixin, View):
    """
    Processamento de importação assistida de contatos a partir de arquivo CSV ou Excel.
    """
    def get(self, request):
        form = ContactUploadImportForm()
        return render(request, 'contacts/import.html', {'form': form})

    def post(self, request):
        form = ContactUploadImportForm(request.POST, request.FILES)
        if not form.is_valid():
            return render(request, 'contacts/import.html', {'form': form})

        uploaded_file = request.FILES['file']
        filename = uploaded_file.name.lower()
        success_count = 0
        error_count = 0
        errors = []

        try:
            if filename.endswith('.csv'):
                decoded = uploaded_file.read().decode('utf-8-sig')
                reader = csv.DictReader(io.StringIO(decoded))
                rows = list(reader)
            elif filename.endswith('.xlsx'):
                wb = openpyxl.load_workbook(uploaded_file, data_only=True)
                ws = wb.active
                headers = [str(cell.value).strip() if cell.value else '' for cell in ws[1]]
                rows = []
                for row in ws.iter_rows(min_row=2, values_only=True):
                    if any(row):
                        rows.append(dict(zip(headers, row)))
            else:
                messages.error(request, _("Formato não suportado. Envie um arquivo .xlsx ou .csv."))
                return render(request, 'contacts/import.html', {'form': form})

            for idx, row in enumerate(rows, start=2):
                name = row.get('nome_razao_social') or row.get('name')
                if not name:
                    continue

                c_type = str(row.get('tipo_contato', 'COMPANY')).strip().upper()
                contact_type = ContactTypeChoices.INDIVIDUAL if 'IND' in c_type else ContactTypeChoices.COMPANY

                doc_type = str(row.get('tipo_documento', 'CNPJ')).strip().upper()
                doc_type_val = DocTypeChoices.CPF if 'CPF' in doc_type else DocTypeChoices.CNPJ

                raw_doc = str(row.get('numero_documento', '') or row.get('doc_number', '')).strip()

                try:
                    contact = Contact(
                        name=name,
                        trade_name=row.get('nome_fantasia', '') or '',
                        contact_type=contact_type,
                        doc_type=doc_type_val,
                        doc_number=raw_doc,
                        email=row.get('email', '') or '',
                        phone=row.get('telefone', '') or '',
                        mobile=row.get('celular', '') or '',
                        city=row.get('cidade', '') or '',
                        state=row.get('estado', '') or '',
                        postal_code=row.get('cep', '') or '',
                    )
                    contact.full_clean()
                    contact.save()
                    success_count += 1
                except Exception as e:
                    error_count += 1
                    errors.append(f"Linha {idx} ({name}): {str(e)}")

            if success_count > 0:
                messages.success(request, _(f"Importação concluída: {success_count} contatos importados com sucesso!"))
            if error_count > 0:
                messages.warning(request, _(f"{error_count} registros não puderam ser importados."))

            return render(request, 'contacts/import.html', {
                'form': form,
                'success_count': success_count,
                'error_count': error_count,
                'errors': errors[:10]
            })

        except Exception as exc:
            messages.error(request, _(f"Erro ao ler o arquivo de dados: {str(exc)}"))
            return render(request, 'contacts/import.html', {'form': form})

import json
import urllib.request
from django.core.cache import cache

class ContactCEPLookupView(LoginRequiredMixin, View):
    """
    Endpoint corporativo para consulta de CEP via API do ViaCEP com cache inteligente em Redis/Memoria
    e tratamento de erros de conexao e timeout.
    """
    def get(self, request):
        raw_cep = request.GET.get('cep', '')
        digits = clean_doc_digits(raw_cep)

        if len(digits) != 8:
            return JsonResponse({
                'found': False,
                'message': _('CEP inválido. O formato deve conter 8 dígitos.')
            }, status=400)

        cache_key = f"viacep_{digits}"
        try:
            cached_data = cache.get(cache_key)
            if cached_data:
                return JsonResponse(cached_data)
        except Exception:
            pass

        url = f"https://viacep.com.br/ws/{digits}/json/"
        req = urllib.request.Request(
            url,
            headers={'User-Agent': 'ERPangea/2.0 (Pangea Engenharia; contatos@pangea.eng.br)'}
        )

        try:
            with urllib.request.urlopen(req, timeout=4.0) as response:
                if response.status != 200:
                    return JsonResponse({
                        'found': False,
                        'message': _('Serviço de CEP indisponível no momento.')
                    }, status=502)
                
                body = response.read().decode('utf-8')
                data = json.loads(body)

                if data.get('erro') in (True, 'true', '1'):
                    return JsonResponse({
                        'found': False,
                        'message': _('CEP não localizado na base dos Correios. Preencha os campos manualmente.')
                    })

                result = {
                    'found': True,
                    'cep': data.get('cep', ''),
                    'street': data.get('logradouro', ''),
                    'complement': data.get('complemento', ''),
                    'neighborhood': data.get('bairro', ''),
                    'city': data.get('localidade', ''),
                    'state': data.get('uf', ''),
                    'country': 'Brasil'
                }

                # Salva em cache por 24 horas (86400s)
                try:
                    cache.set(cache_key, result, timeout=86400)
                except Exception:
                    pass
                return JsonResponse(result)

        except urllib.error.URLError:
            return JsonResponse({
                'found': False,
                'error': True,
                'message': _('Falha de conexão com o serviço de CEP. Você pode preencher os campos manualmente.')
            }, status=503)
        except Exception as e:
            return JsonResponse({
                'found': False,
                'error': True,
                'message': _('Não foi possível obter os dados do CEP no momento. Preencha manualmente.')
            }, status=500)
